# Usage: python scripts/build-rca.py src/test-cases.ts "<output.xlsx>" <UAT results.json> <UAT-internal results.json> <Pre-UAT results.json> "<UAT run>" "<UAT-internal run>" "<Pre-UAT run>"
# Results files: test-results*/results.json, or results-merged.json after re-runs (UAT 13 Oct: .rca-inputs/uat-final-run-13oct.json).
# Builds "Cinescape CMS UAT, UAT internal and Pre-UAT RCA - 14 Oct 2026.xlsx" in the layout of the 14 Oct Web /
# Android RCA workbooks: summary, "All test cases" (result / issue type / reason per environment + remarks) and
# "Failed cases" (with a manual check status per environment).
import json, re, sys
from rca_rules import load_cases, load, reason
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles.differential import DifferentialStyle

cases_file, out_file = sys.argv[1], sys.argv[2]
import os
RCA_DATE = os.environ.get('RCA_DATE', '14 Oct 2026')  # date in the sheet titles
envs = [  # (label, results json, address, run description)
    ('UAT', sys.argv[3], 'http://10.35.1.66:8092', sys.argv[6]),
    ('UAT (internal)', sys.argv[4], 'http://10.35.1.66:8102', sys.argv[7]),
    ('Pre-UAT', sys.argv[5], 'http://10.35.1.66:8122', sys.argv[8]),
]

cases = load_cases(cases_file)

results = [load(e[1]) for e in envs]


thin = Side(style='thin', color='FFBFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
grey = PatternFill('solid', fgColor='FFF2F2F2')
blue = PatternFill('solid', fgColor='FFDDEBF7')
bold = Font(name='Calibri', size=11, bold=True)
body = Font(name='Calibri', size=11)
wrap = Alignment(wrap_text=True, vertical='top')
good = DifferentialStyle(font=Font(color='FF006100'), fill=PatternFill(bgColor='FFC6EFCE'))
bad = DifferentialStyle(font=Font(color='FF9C0006'), fill=PatternFill(bgColor='FFFFC7CE'))
neutral = DifferentialStyle(font=Font(color='FF9C5700'), fill=PatternFill(bgColor='FFFFEB9C'))

def cf(ws, col, first, last, kind):
    from openpyxl.formatting.rule import Rule
    rng = f'{col}{first}:{col}{last}'
    pairs = [('Passed', good), ('Failed', bad)] if kind == 'result' else [('Application issue', bad), ('Script issue', neutral)]
    for value, style in pairs:
        rule = Rule(type='expression', dxf=style, formula=[f'${col}{first}="{value}"'])
        ws.conditional_formatting.add(rng, rule)

def header(ws, row, values, fill=grey):
    for i, v in enumerate(values, 1):
        if v is None: continue
        c = ws.cell(row, i, v); c.font, c.fill, c.border, c.alignment = bold, fill, border, wrap

def set_cell(ws, row, col, v):
    c = ws.cell(row, col, v if v != '' else None); c.font, c.border, c.alignment = body, border, wrap

wb = openpyxl.Workbook()
# ---------------- All test cases ----------------
ws = wb.active; ws.title = 'All test cases'
ws['A1'] = f'All test cases - UAT, UAT (internal) and Pre-UAT ({RCA_DATE})'; ws['A1'].font = Font(name='Calibri', size=13, bold=True)
cols = ['Test ID', 'Test case', 'Description (expected result)', 'Test steps']
for label, *_ in envs: cols += [f'{label} result', f'{label} issue type', f'{label} reason']
cols.append('Remarks')
header(ws, 3, cols)
for i, w in enumerate([12, 30, 36, 55] + [11, 16, 24] * 3 + [80], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
failed_rows = []
row = 4
for c in cases:
    cid = c['id']
    title = f"{c['menu']} > {c['submenu']} - {c['feature']}" if c['submenu'] else f"{c['menu']} - {c['feature']}"
    vals = [cid, title, c['expected'], c['steps']]
    remarks, per_env = [], []
    for (label, *_), res in zip(envs, results):
        r = res.get(cid, {'status': 'Not run', 'error': '', 'notes': [], 'step': ''})
        if r['status'] == 'Failed':
            issue, short, text = reason(cid, r, label)
            remarks.append((label, text))
        else:
            issue, short = '', ''
            for n in r['notes']:
                if n['type'] == 'skip': remarks.append((label, f"Not run: {n['description']}"))
                elif n['type'] == 'finding': remarks.append((label, f"Finding: {n['description']}"))
        vals += [r['status'], issue, short]
        per_env.append((r['status'], issue, short))
    # Same remark on every environment: written once, without the environment name.
    texts = {}
    for label, t in remarks: texts.setdefault(t, []).append(label)
    lines = [t if len(labels) == len(envs) else f"{', '.join(labels)}: {t}" for t, labels in texts.items()]
    if cid == 'CMS-TC-177' and per_env[0][0] == 'Failed':
        lines.append('UAT: intermittent - the same form saved banners in CMS-TC-178/182/183 of the same run, and CMS-TC-177 passed when run again alone.')
    vals.append('\n'.join(lines))
    for i, v in enumerate(vals, 1): set_cell(ws, row, i, v)
    if any(p[0] == 'Failed' for p in per_env): failed_rows.append((cid, title, per_env, '\n'.join(lines)))
    row += 1
last = row - 1
ws.freeze_panes = 'C4'; ws.auto_filter.ref = f'A3:{openpyxl.utils.get_column_letter(len(cols))}{last}'
dv = DataValidation(type='list', formula1='"Application issue,Script issue"', allow_blank=True)
ws.add_data_validation(dv)
for k in range(len(envs)):
    rc = openpyxl.utils.get_column_letter(5 + 3 * k); ic = openpyxl.utils.get_column_letter(6 + 3 * k)
    cf(ws, rc, 4, last, 'result'); cf(ws, ic, 4, last, 'issue'); dv.add(f'{ic}4:{ic}{last}')

# ---------------- Failed cases ----------------
fs = wb.create_sheet('Failed cases')
fs['A1'] = f'Failed cases - UAT, UAT (internal) and Pre-UAT ({RCA_DATE})'; fs['A1'].font = Font(name='Calibri', size=13, bold=True)
fs['A2'] = 'Manual check status: choose Passed or Failed after checking the case by hand on that environment.'; fs['A2'].font = body
fcols = ['Test ID', 'Test case']
for label, *_ in envs: fcols += [f'{label} result', f'{label} issue type', f'{label} reason']
fcols += ['Remarks', 'Manual check status', None, None]
header(fs, 3, fcols)
mc = len(fcols) - 2  # first manual check column
header(fs, 4, [None] * (mc - 1) + [label for label, *_ in envs], blue)
fs.cell(3, mc).fill = blue
for col in range(1, mc - 1 + 1):
    fs.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
fs.merge_cells(start_row=3, start_column=mc, end_row=3, end_column=mc + 2)
for i, w in enumerate([12, 30] + [11, 16, 24] * 3 + [70, 14, 14, 14], 1):
    fs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
r = 5
for cid, title, per_env, remark in failed_rows:
    vals = [cid, title] + [x for p in per_env for x in p] + [remark, '', '', '']
    for i, v in enumerate(vals, 1): set_cell(fs, r, i, v)
    r += 1
flast = max(r - 1, 5)
fs.freeze_panes = 'C5'; fs.auto_filter.ref = f'A4:{openpyxl.utils.get_column_letter(mc + 2)}{flast}'
dv2 = DataValidation(type='list', formula1='"Passed,Failed"', allow_blank=True); fs.add_data_validation(dv2)
for k in range(len(envs)):
    cf(fs, openpyxl.utils.get_column_letter(3 + 3 * k), 5, flast, 'result'); cf(fs, openpyxl.utils.get_column_letter(4 + 3 * k), 5, flast, 'issue')
for k in range(3):
    col = openpyxl.utils.get_column_letter(mc + k); cf(fs, col, 5, flast, 'result'); dv2.add(f'{col}5:{col}{flast}')

# ---------------- Summary (first sheet) ----------------
sm = wb.create_sheet('RCA Report', 0)
sm['A1'] = 'Cinescape CMS - RCA report (UAT, UAT internal and Pre-UAT)'; sm['A1'].font = Font(name='Calibri', size=13, bold=True)
notes = [
    'Run date: ' + '; '.join(f'{label} {run}' for label, _, _, run in envs) + ' (Playwright automated cases CMS-TC-001 to CMS-TC-250)',
    '     '.join(f'{label}: {url}' for label, _, url, _ in envs),
    'Issue type: "Application issue" = the CMS needs a fix.  "Script issue" = the CMS changed and the test needs an update. Problems in the tests themselves were found in earlier runs and are fixed: no Script issue is open.',
    'Reason: the short cause of the failure. The full details of every test case (description, steps, results and remarks) are on the sheet "All test cases"; the failed ones, with a manual check status, on "Failed cases".',
    '"Not run": no data on that environment in the last 30 days, no Reset button on the report (manual checklist: NA), or deferred in the sheet (CMS-TC-044).',
]
for i, t in enumerate(notes, 2): sm.cell(i, 1, t).font = body
header(sm, 8, [None] + [label for label, *_ in envs])
sm.column_dimensions['A'].width = 34
for k in range(len(envs)): sm.column_dimensions[openpyxl.utils.get_column_letter(2 + k)].width = 16
labels = ['Cases run', 'Passed', 'Failed', '  Failed - Application issue', '  Failed - Script issue', 'Not run', 'Pass rate']
for i, lab in enumerate(labels):
    rr = 9 + i
    a = sm.cell(rr, 1, lab); a.font, a.border = body, border
    for k in range(len(envs)):
        res = f"'All test cases'!${openpyxl.utils.get_column_letter(5 + 3 * k)}$4:${openpyxl.utils.get_column_letter(5 + 3 * k)}${last}"
        iss = f"'All test cases'!${openpyxl.utils.get_column_letter(6 + 3 * k)}$4:${openpyxl.utils.get_column_letter(6 + 3 * k)}${last}"
        col = openpyxl.utils.get_column_letter(2 + k)
        f = {
            'Cases run': f"=COUNTA('All test cases'!$A$4:$A${last})-COUNTIF({res},\"Not run\")",
            'Passed': f'=COUNTIF({res},"Passed")', 'Failed': f'=COUNTIF({res},"Failed")',
            '  Failed - Application issue': f'=COUNTIFS({res},"Failed",{iss},"Application issue")',
            '  Failed - Script issue': f'=COUNTIFS({res},"Failed",{iss},"Script issue")',
            'Not run': f'=COUNTIF({res},"Not run")', 'Pass rate': f'=IFERROR({col}10/{col}9,0)',
        }[lab]
        c = sm.cell(rr, 2 + k, f); c.font, c.border = body, border
        if lab == 'Pass rate': c.number_format = '0%'

# Failures by reason, per environment
rr = 18
sm.cell(rr, 1, 'Failures by reason (Application issue)').font = bold
header(sm, rr + 1, ['Reason'] + [label for label, *_ in envs] + ['Test IDs'])
sm.column_dimensions[openpyxl.utils.get_column_letter(2 + len(envs))].width = 60
groups = {}
for cid, title, per_env, _ in failed_rows:
    for k, (st, issue, short) in enumerate(per_env):
        if st == 'Failed': groups.setdefault(short, [set() for _ in envs])[k].add(cid)
rr += 2
for short, sets in sorted(groups.items(), key=lambda kv: -sum(len(s) for s in kv[1])):
    allids = sorted(set().union(*sets))
    vals = [short] + [len(s) or None for s in sets] + [', '.join(x[-3:] for x in allids)]
    for i, v in enumerate(vals, 1): set_cell(sm, rr, i, v)
    rr += 1
rr += 1
sm.cell(rr, 1, 'Findings (the cases pass; noted for the developers)').font = bold
for t in [
    'Delete and Re-send Password raise a script error ("Cannot set properties of null (setting \'href\')": #delkey is not on the page); the action itself works.',
    'A Payment Mode name over 30 characters gives "Something went wrong, Please try again." instead of a clear message or a length limit.',
    'Food and Beverages, Cancellations and Refund reports start with empty dates; Submit then answers "Invalid date format." instead of asking for the dates.',
    'UAT: test data that cannot be removed through the CMS (Genre has no delete): genre "AUTOTEST T040-1011101157" (id CMS-897C62B5) - please remove it in the database.',
    'UAT (internal) and Pre-UAT share the same data (a record changed on one shows on the other); UAT has its own. Their test runs are done one after the other.',
    'UAT (internal) runs the older CMS pages (Login button, no Light / Dark switch, Summernote editors, Add Food still offered); the tests handle both designs.',
]:
    rr += 1; sm.cell(rr, 1, t).font = body
# Page script errors of cases that work (reported as findings since 21 Oct 2026), with their cases and environments.
script_errors = {}
for (label, *_), res in zip(envs, results):
    for cid, r in res.items():
        for n in r['notes']:
            if n['type'] == 'finding' and n['description'].startswith('Page script error:'):
                script_errors.setdefault(n['description'], {}).setdefault(cid, []).append(label)
for text, by_case in sorted(script_errors.items(), key=lambda kv: min(kv[1])):
    labels = sorted({l for ls in by_case.values() for l in ls}, key=[e[0] for e in envs].index)
    where = '' if len(labels) == len(envs) else f" ({', '.join(labels)})"
    rr += 1; sm.cell(rr, 1, f"{text} - cases {', '.join(c[-3:] for c in sorted(by_case))}{where}").font = body

wb.save(out_file)
print(out_file)
for (label, *_), res in zip(envs, results):
    print(label, {k: sum(1 for v in res.values() if v['status'] == k) for k in ('Passed', 'Failed', 'Not run')}, 'missing', 250 - len(res))
print('failed cases sheet rows:', len(failed_rows), '| reasons:', len(groups), '| See remarks:', [c for c, _, p, _ in failed_rows if any(x[2] == 'See remarks' for x in p)])
