# Usage: python scripts/build-status.py src/test-cases.ts "<output.xlsx>" <UAT results.json> <UAT-internal results.json> <Pre-UAT results.json>
# Builds the test execution status workbook: a summary, one sheet per environment with every test case (steps,
# expected result, latest Passed / Failed / Not run status, reason, remarks, when it ran) and a side-by-side sheet.
# STATUS_DATE sets the date in the titles.
import os, sys
from datetime import datetime
import openpyxl
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as col_letter
from rca_rules import load_cases, load, reason

cases_file, out_file = sys.argv[1], sys.argv[2]
DATE = os.environ.get('STATUS_DATE', datetime.now().strftime('%d %b %Y'))
envs = [  # (sheet name, results json, address)
    ('UAT', sys.argv[3], 'http://10.35.1.66:8092'),
    ('UAT (internal)', sys.argv[4], 'http://10.35.1.66:8102'),
    ('Pre-UAT', sys.argv[5], 'http://10.35.1.66:8122'),
]
cases = load_cases(cases_file)
results = [load(path) for _, path, _ in envs]

thin = Side(style='thin', color='FFBFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='FF1F4E78')
head_font = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
title_font = Font(name='Calibri', size=14, bold=True, color='FF1F4E78')
bold = Font(name='Calibri', size=11, bold=True)
body = Font(name='Calibri', size=11)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top', wrap_text=True)
styles = {
    'Passed': DifferentialStyle(font=Font(color='FF006100', bold=True), fill=PatternFill(bgColor='FFC6EFCE')),
    'Failed': DifferentialStyle(font=Font(color='FF9C0006', bold=True), fill=PatternFill(bgColor='FFFFC7CE')),
    'Not run': DifferentialStyle(font=Font(color='FF7F6000'), fill=PatternFill(bgColor='FFFFEB9C')),
}


def header(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font, c.fill, c.border, c.alignment = head_font, head_fill, border, center
    ws.row_dimensions[row].height = 30


def put(ws, row, col, v, align=wrap):
    c = ws.cell(row, col, None if v in ('', None) else v)
    c.font, c.border, c.alignment = body, border, align
    return c


def colour_status(ws, col, first, last):
    for value, style in styles.items():
        ws.conditional_formatting.add(f'{col}{first}:{col}{last}',
                                      Rule(type='expression', dxf=style, formula=[f'${col}{first}="{value}"']))


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[col_letter(i)].width = w


def ran_at(r):
    if not r.get('start'):
        return None
    return datetime.fromisoformat(r['start'].replace('Z', '+00:00')).astimezone().replace(tzinfo=None)


def detail(cid, r, env):
    """Issue type, reason and remarks of one case on one environment."""
    if r['status'] == 'Failed':
        issue, short, text = reason(cid, r, env)
        return issue, short, text
    notes = [('Not run: ' if n['type'] == 'skip' else 'Finding: ') + n['description'] for n in r['notes']]
    return '', '', '\n'.join(notes)


def title_of(c):
    return f"{c['menu']} > {c['submenu']} - {c['feature']}" if c['submenu'] else f"{c['menu']} - {c['feature']}"


wb = openpyxl.Workbook()
FIRST = 4  # first data row of the environment sheets
last = FIRST + len(cases) - 1
STATUS_COL = 'H'

# ---------------- One sheet per environment ----------------
per_env = []  # [(status, reason)] per case, per environment
for (name, _, url), res in zip(envs, results):
    ws = wb.create_sheet(name)
    ws['A1'] = f'Cinescape CMS - {name} ({url}) - test execution status, {DATE}'
    ws['A1'].font = title_font
    ws['A2'] = 'Latest result of every automated test case on this environment. "Executed on" shows when the result was produced.'
    ws['A2'].font = body
    cols = ['S.No', 'Test ID', 'Menu', 'Sub menu', 'Feature', 'Test steps', 'Expected result', 'Status',
            'Issue type', 'Reason', 'Failed step', 'Remarks', 'Executed on', 'Duration (s)', 'Spec file']
    header(ws, 3, cols)
    widths(ws, [6, 12, 18, 20, 26, 60, 50, 11, 17, 28, 34, 70, 17, 11, 52])
    statuses = []
    for i, c in enumerate(cases):
        row = FIRST + i
        r = res.get(c['id'], {'status': 'Not run', 'error': '', 'notes': [], 'step': ''})
        issue, short, remarks = detail(c['id'], r, name)
        vals = [i + 1, c['id'], c['menu'], c['submenu'], c['feature'], c['steps'], c['expected'], r['status'],
                issue, short, r.get('step', '') if r['status'] == 'Failed' else '', remarks]
        for k, v in enumerate(vals, 1):
            put(ws, row, k, v, center if k in (1, 8) else wrap)
        put(ws, row, 13, ran_at(r), center).number_format = 'dd-mmm-yyyy hh:mm'
        put(ws, row, 14, round(r['duration'] / 1000, 1) if r.get('duration') else None, center)
        put(ws, row, 15, r.get('file', ''))
        statuses.append((r['status'], short))
    per_env.append(statuses)
    colour_status(ws, STATUS_COL, FIRST, last)
    ws.freeze_panes = 'C4'
    ws.auto_filter.ref = f'A3:{col_letter(len(cols))}{last}'

# ---------------- All environments side by side ----------------
ae = wb.create_sheet('All environments')
ae['A1'] = f'Cinescape CMS - status on all environments, {DATE}'
ae['A1'].font = title_font
ae['A2'] = 'Filter a Status column to "Failed" to list the failures of one environment; the reasons are on the environment sheets.'
ae['A2'].font = body
cols = ['S.No', 'Test ID', 'Menu', 'Sub menu', 'Feature']
for name, *_ in envs:
    cols += [f'{name} status', f'{name} reason']
cols.append('Same on all')
header(ae, 3, cols)
widths(ae, [6, 12, 18, 20, 30] + [12, 26] * len(envs) + [11])
for i, c in enumerate(cases):
    row = FIRST + i
    vals = [i + 1, c['id'], c['menu'], c['submenu'], c['feature']]
    for statuses in per_env:
        vals += list(statuses[i])
    vals.append('Yes' if len({s[i][0] for s in per_env}) == 1 else 'No')
    for k, v in enumerate(vals, 1):
        put(ae, row, k, v, center if k == 1 or (k > 5 and (k - 6) % 2 == 0) or k == len(vals) else wrap)
for k in range(len(envs)):
    colour_status(ae, col_letter(6 + 2 * k), FIRST, last)
ae.freeze_panes = 'C4'
ae.auto_filter.ref = f'A3:{col_letter(len(cols))}{last}'

# ---------------- Summary (first sheet) ----------------
sm = wb.create_sheet('Summary', 0)
sm['A1'] = f'Cinescape CMS - Test execution status ({DATE})'
sm['A1'].font = title_font
sm['A2'] = f'{len(cases)} automated Playwright test cases (CMS-TC-001 to CMS-TC-{len(cases):03d}), run on each environment below.'
sm['A2'].font = body
header(sm, 4, ['Environment', 'Address', 'Executed between', 'Total', 'Passed', 'Failed', 'Not run', 'Pass rate (of run)'])
widths(sm, [26, 26, 34, 9, 9, 9, 9, 16])
for k, ((name, _, url), res) in enumerate(zip(envs, results)):
    row = 5 + k
    times = sorted(t for t in (ran_at(r) for r in res.values()) if t)
    span = f"{times[0]:%d %b %Y} - {times[-1]:%d %b %Y}" if times and times[0].date() != times[-1].date() else (f'{times[0]:%d %b %Y}' if times else '')
    rng = f"'{name}'!${STATUS_COL}${FIRST}:${STATUS_COL}${last}"
    vals = [name, url, span, f"=COUNTA('{name}'!$B${FIRST}:$B${last})", f'=COUNTIF({rng},"Passed")',
            f'=COUNTIF({rng},"Failed")', f'=COUNTIF({rng},"Not run")', f'=IFERROR(E{row}/(E{row}+F{row}),0)']
    for i, v in enumerate(vals, 1):
        put(sm, row, i, v, wrap if i <= 3 else center)
    sm.cell(row, 8).number_format = '0.0%'

# Results by menu
row = 5 + len(envs) + 2
sm.cell(row, 1, 'Results by menu').font = bold
row += 1
cols = ['Menu', 'Test cases']
for name, *_ in envs:
    cols += [f'{name} passed', f'{name} failed', f'{name} not run']
header(sm, row, cols)
for k in range(len(envs) * 3):
    sm.column_dimensions[col_letter(3 + k)].width = max(sm.column_dimensions[col_letter(3 + k)].width or 0, 12)
menus = list(dict.fromkeys(c['menu'] for c in cases))
for m in menus:
    row += 1
    vals = [m, f"=COUNTIF('{envs[0][0]}'!$C${FIRST}:$C${last},A{row})"]
    for name, *_ in envs:
        for st in ('Passed', 'Failed', 'Not run'):
            vals.append(f"=COUNTIFS('{name}'!$C${FIRST}:$C${last},$A{row},'{name}'!${STATUS_COL}${FIRST}:${STATUS_COL}${last},\"{st}\")")
    for i, v in enumerate(vals, 1):
        put(sm, row, i, v, wrap if i == 1 else center)

row += 2
sm.cell(row, 1, 'Notes').font = bold
for t in [
    'Passed: every step and check of the case passed.  Failed: a check did not hold (the reason and remarks are on the environment sheet).  '
    'Not run: no data on that environment in the last 30 days, no Reset button on the report, or deferred (CMS-TC-044).',
    'Every open failure is an Application issue (the CMS needs a fix); no Script issue is open.',
    'Each case has a screenshot per step and a video in the Playwright report and dashboard of its environment.',
]:
    row += 1
    sm.cell(row, 1, t).font = body

del wb['Sheet']
wb.save(out_file)
print(out_file)
for (name, *_), statuses in zip(envs, per_env):
    print(name, {s: sum(1 for x in statuses if x[0] == s) for s in ('Passed', 'Failed', 'Not run')})
