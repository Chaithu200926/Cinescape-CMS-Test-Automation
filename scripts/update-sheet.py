# Rewrites Test Steps / Expected Result of CMS_Test_Cases.xlsx from the automated steps (next/test-steps.json).
# Usage: python scripts/update-sheet.py "<CMS_Test_Cases.xlsx>" next/test-steps.json scripts/data/checklist-remarks.json src/cms-forms.ts tests
# Rewrites Test Steps / Expected Result of CMS_Test_Cases.xlsx from the automated tests' steps (1:1), adds the
# manual-checklist remark and the automated test file, and adds CMS-TC-245..250 (Kiosk Transaction Report).
import sys, json, re, copy, os, glob
import openpyxl

xlsx, steps_file, checklist_file, forms_file, tests_dir = sys.argv[1:6]
steps = json.load(open(steps_file, encoding='utf8'))
checklist = json.load(open(checklist_file, encoding='utf8'))
forms = open(forms_file, encoding='utf8').read()
spec_files = {os.path.basename(f)[:10].upper(): os.path.basename(f) for f in glob.glob(os.path.join(tests_dir, 'cms-tc-*.spec.ts'))}

# Field each page's Edit case changes (from src/cms-forms.ts), for readable step text.
edit_field = {}
for m in re.finditer(r"'([^']+ › [^']+)':\s*(?:tabForm\(|\{)([\s\S]*?)(?=\n  '[^']+ › |\n\};)", forms):
    f = re.search(r"edit:\s*\{\s*field:\s*'(\w+)'", m.group(2))
    edit_field[m.group(1)] = f.group(1) if f else 'name'
for tab in ['Configurations › Ratings', 'Configurations › T & C', 'Configurations › Privacy Policy']:
    edit_field.setdefault(tab, 'name')

def readable(title, page):
    field = edit_field.get(page, 'the field')
    subs = [
        (r'\$\{record\.name\}', 'AUTOTEST T<case>-<time>'), (r'\$\{value\}', '<value from the first row>'),
        (r'\$\{column\.name\}', '<first sortable column>'), (r'\$\{Math\.min\(size, total\)\}', 'N (or all, if fewer)'),
        (r'\$\{size\}', 'N'), (r'\$\{to - from \+ 1\}', 'y - x + 1'), (r'\$\{from\}', 'x'), (r'\$\{to\}', 'y'),
        (r'All \$\{total\} entries fit on one page: check', 'If all entries fit on one page: check'),
        (r'\$\{lastRow \+ 1\}', 'y + 1'), (r'\$\{field\}', field), (r'\$\{newValue\}', '<value> EDITED'),
        (r'\$\{changed\}', '<current value> AUTOTEST'), (r'\$\{original\}', '<original value>'), (r'\$\{foodId\}', '<food id>'),
        (r'\$\{cinema\}', '<first cinema>'), (r'\$\{id\}', '<transaction id>'), (r'\$\{role\}', 'Report_Admin'),
        (r'\$\{badEmail\}', 'autotest-<token>-no-at-sign'), (r'\$\{file\.name\}', '<file>'), (r'\$\{table\.rows\.length\}', 'N'),
        (r'\$\{other\}', 'the other theme'), (r'\$\{start\}', 'the original theme'), (r'\$\{value\}', '<value>'),
    ]
    for a, b in subs:
        title = re.sub(a, b, title)
    return re.sub(r'\$\{[^}]*\}', '<…>', title)

wb = openpyxl.load_workbook(xlsx)
ws = wb['Test Cases']
header_style = ws['H1']
body_style = ws['E2']

def style_like(cell, src):
    cell.font = copy.copy(src.font); cell.fill = copy.copy(src.fill); cell.border = copy.copy(src.border)
    cell.alignment = copy.copy(src.alignment); cell.number_format = src.number_format

for col, title, width in [('I', 'Manual Checklist (cms test.xlsx)', 48), ('J', 'Automated Test (tests/)', 46)]:
    ws[f'{col}1'] = title
    style_like(ws[f'{col}1'], header_style)
    ws.column_dimensions[col].width = width
ws.column_dimensions['F'].width = 56

rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
kiosk = {245: 'List', 246: 'Sort', 247: 'Search', 248: 'Download', 249: 'Pagination', 250: 'Reset'}
for n, feature in kiosk.items():
    cid = f'CMS-TC-{n}'
    if cid in rows: continue
    r = ws.max_row + 1
    for col, v in zip('ABCDH', ['', 'Reports', 'Kiosk Transaction Report', feature, 'Not Run']):
        ws[f'{col}{r}'] = v
    ws[f'A{r}'] = cid
    for col in 'ABCDEFGHIJ':
        style_like(ws[f'{col}{r}'], body_style)
    rows[cid] = r

updated = 0
for cid, r in rows.items():
    info = steps.get(cid)
    if not info: continue
    menu, sub = ws.cell(r, 2).value or '', ws.cell(r, 3).value or ''
    page = f'{menu} › {sub}'
    if info.get('deferred'):
        ws[f'E{r}'] = 'Deferred: functionality not completed (will be tested later).'
        ws[f'F{r}'] = '-'
    else:
        lines = [readable(s, page) for s in info['steps']]
        checks = [re.sub(r'^Check (that )?', '', l) for l in lines if re.match(r'^(Check|If all)', l)]
        ws[f'E{r}'] = '\n'.join(f'{i}. {l}' for i, l in enumerate(lines, 1))
        ws[f'F{r}'] = '\n'.join(f'{i}. {l[0].upper() + l[1:]}' for i, l in enumerate(checks, 1)) or 'The action completes without an error.'
    ws[f'I{r}'] = checklist.get(cid, '')
    ws[f'J{r}'] = spec_files.get(cid, '')
    for col in 'EFIJ':
        style_like(ws[f'{col}{r}'], body_style)
    updated += 1

ws.auto_filter.ref = f'A1:J{ws.max_row}'
wb.save(xlsx)
print('updated', updated, 'rows; total rows', ws.max_row - 1)
